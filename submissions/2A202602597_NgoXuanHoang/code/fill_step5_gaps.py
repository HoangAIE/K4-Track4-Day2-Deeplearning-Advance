# -*- coding: utf-8 -*-
"""Buoc 5: (1) ve curves per-seed cho F01 tu history.csv; (2) dien sheet Summary top-10 theo val macro-F1."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT_DIR = Path(__file__).resolve().parent.parent
CODE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT_DIR))
sys.path.insert(0, str(CODE_DIR))

from train import plot_curves

# ---------- 1. F01 per-seed curves ----------
for seed in [0, 1, 2]:
    hist = pd.read_csv(ROOT_DIR / f"runs/F01/seed{seed}/history.csv").to_dict("records")
    out = ROOT_DIR / f"curves/F01_seed{seed}_convnext_tiny.png"
    plot_curves(hist, out,
                f"F01: convnext_tiny CutMix (seed={seed}, finetune, CE, 12ep)")
    print(f"curve seed{seed}: best F1={max(h['val_macro_f1'] for h in hist)*100:.2f}% -> {out.name}")

# ---------- 2. Summary top-10 ----------
wb = openpyxl.load_workbook(ROOT_DIR / "results.xlsx")
ws = wb["Summary"]
if ws.max_row > 1:
    ws.delete_rows(2, ws.max_row - 1)

HFILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HFONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DFONT = Font(name="Calibri", size=10)
BFILL = PatternFill(start_color="E8F8F5", end_color="E8F8F5", fill_type="solid")
THIN = Border(left=openpyxl.styles.Side(style="thin", color="D9D9D9"),
              right=openpyxl.styles.Side(style="thin", color="D9D9D9"),
              top=openpyxl.styles.Side(style="thin", color="D9D9D9"),
              bottom=openpyxl.styles.Side(style="thin", color="D9D9D9"))
for i in range(1, 11):
    c = ws.cell(row=1, column=i)
    c.fill = HFILL
    c.font = HFONT
    c.alignment = HALIGN
    c.border = THIN

# Thu thap ung vien: Backbones (B) + Training (T) + F01 val mean + Inference best
cands = []
for exp in ["B01", "B02", "B03", "B04", "B05"]:
    s = json.load(open(ROOT_DIR / f"runs/{exp}/seed0/summary.json", encoding="utf-8"))
    cands.append({"exp_id": exp, "backbone": s["backbone"], "cfg": f"T00 recipe ({exp})",
                  "f1": s["best_val_macro_f1"], "top1": s["best_val_top1"],
                  "lat": None, "params": s["params_m"], "gmacs": s["gmacs"],
                  "note": "Buoc 1: so sanh backbone"})
import openpyxl as _oxl
wb0 = _oxl.load_workbook(ROOT_DIR / "results.xlsx", data_only=True)
wtr = wb0["Training"]
for row in wtr.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    cands.append({"exp_id": row[0], "backbone": row[1], "cfg": str(row[3])[:60],
                  "f1": float(row[5]), "top1": float(row[6]),
                  "lat": None, "params": 27.827, "gmacs": 4.455,
                  "note": f"Buoc 2: {row[2]}"})
winf = wb0["Inference"]
inf_best = {}
for row in winf.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    inf_best[row[0]] = (float(row[4]), float(row[5]),
                        row[7] if isinstance(row[7], (int, float)) else None)
for iid, (f1, top1, lat) in inf_best.items():
    if iid in ("I05a", "I06", "I01"):
        cands.append({"exp_id": iid, "backbone": "convnext_tiny",
                      "cfg": f"Buoc 3 tren T04 ({iid})", "f1": f1, "top1": top1,
                      "lat": lat, "params": 27.827, "gmacs": 4.455,
                      "note": "Buoc 3: suy luan (val)"})
# F01 val mean (3 seed)
f1s = [json.load(open(ROOT_DIR / f"runs/F01/seed{s}/summary.json", encoding="utf-8"))["best_val_macro_f1"]
       for s in [0, 1, 2]]
cands.append({"exp_id": "F01(val)", "backbone": "convnext_tiny",
              "cfg": "Chung ket: CutMix+T-scale (mean 3 seed, val)",
              "f1": round(sum(f1s) / 3, 4), "top1": None, "lat": 3.45,
              "params": 27.827, "gmacs": 4.455, "note": "Buoc 4: val mean (test 0.9734)"})

cands.sort(key=lambda d: d["f1"], reverse=True)
LAT_B = {"B01": 12.47, "B02": 11.44, "B03": 19.56, "B04": 15.67, "B05": 12.80}
r = 2
for rank, d in enumerate(cands[:10], start=1):
    lat = d["lat"] if d["lat"] is not None else LAT_B.get(d["exp_id"], "")
    if d["exp_id"].startswith("T") and d["exp_id"] != "T00":
        lat = 3.45  # cung backbone convnext_tiny FP32 batch-1 (do Buoc 3)
    elif d["exp_id"] == "T00":
        lat = 3.45
    vals = [rank, d["exp_id"], d["backbone"], d["cfg"], round(d["f1"], 4),
            round(d["top1"], 4) if d["top1"] is not None else "", lat,
            d["params"], d["gmacs"], d["note"]]
    for ci, v in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=ci, value=v)
        cell.font = DFONT
        cell.border = THIN
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if ci in (5, 6):
            cell.number_format = "0.0000"
        if ci in (7, 8, 9):
            cell.number_format = "0.00"
    if rank == 1:
        for ci in range(1, 11):
            ws.cell(row=r, column=ci).fill = BFILL
    r += 1
for i, w in enumerate([8, 10, 26, 44, 13, 12, 16, 12, 10, 40], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
wb.save(ROOT_DIR / "results.xlsx")
print("Summary top-10 written. Sheets:", wb.sheetnames)
