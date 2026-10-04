from pathlib import Path
import openpyxl

repo_root = Path("d:/K4-Track4-Day2-Deeplearning-Advance")

# 1. Verify results.xlsx
excel_p = repo_root / "results.xlsx"
assert excel_p.exists(), "results.xlsx missing"

wb = openpyxl.load_workbook(excel_p)
assert "Backbones" in wb.sheetnames, "Backbones sheet missing"
ws = wb["Backbones"]

expected_headers = [
    "exp_id", "backbone", "tag trọng số", "#tham số (M)", "GMAC", 
    "độ phân giải", "epoch", "seed", "macro-F1 val", "top-1 val", 
    "thời gian train/epoch", "độ trễ batch-1 (ms)", "ghi chú"
]

actual_headers = [ws.cell(row=1, column=c).value for c in range(1, 14)]
assert actual_headers == expected_headers, f"Header mismatch: {actual_headers} vs {expected_headers}"
assert ws.freeze_panes == "A2", f"Freeze panes: {ws.freeze_panes}"
assert ws.max_row == 6, f"Expected 6 rows (header + 5 backbones), got {ws.max_row}"

rows = []
for r in range(2, 7):
    row_data = [ws.cell(row=r, column=c).value for c in range(1, 14)]
    rows.append(row_data)
    print(f"Row {r}: {row_data[0]} | {row_data[1]} | F1={row_data[8]} | Top1={row_data[9]} | Latency={row_data[11]} | Fill={ws.cell(row=r, column=1).fill.start_color.rgb}")

# Row 3 (B02: convnext_tiny) must have highlight fill
fill_color = ws.cell(row=3, column=1).fill.start_color.rgb
print(f"B02 highlight fill color: {fill_color}")

# 2. Verify curves
curve_files = [
    "B01_resnet50.png",
    "B02_convnext_tiny.png",
    "B03_swin_tiny_patch4_window7_224.png",
    "B04_efficientnet_b0.png",
    "B05_mobilenetv3_large_100.png"
]

print("\n--- Verifying curves ---")
for cf in curve_files:
    p = repo_root / "curves" / cf
    assert p.exists(), f"Curve {cf} missing"
    size_kb = p.stat().st_size / 1024
    print(f"Curve {cf}: size = {size_kb:.1f} KB")
    assert size_kb > 20, f"Curve {cf} is too small ({size_kb} KB)"

print("\nEXCEL AND CURVES VERIFICATION PASSED!")
