import json
import openpyxl
import torch
import numpy as np
import pandas as pd
from pathlib import Path
import sys
sys.path.insert(0, str(Path('.').resolve()))
import eval as ev

wb = openpyxl.load_workbook('results.xlsx')
ws = wb['Backbones']

with open('runs/benchmark_m1.json', 'r', encoding='utf-8') as f:
    bench = json.load(f)

print('--- AUDITING RUN ARTIFACTS AND METRIC INTEGRITY ---')
val_df = pd.read_csv('data/labels/val_subset0.csv')
y_true_expected = val_df['Label'].values

for r in range(2, 7):
    exp_id = ws.cell(r, 1).value
    backbone = ws.cell(r, 2).value
    tag = ws.cell(r, 3).value
    params_excel = float(ws.cell(r, 4).value)
    gmac_excel = float(ws.cell(r, 5).value)
    res_excel = int(ws.cell(r, 6).value)
    epoch_excel = int(ws.cell(r, 7).value)
    seed_excel = int(ws.cell(r, 8).value)
    f1_excel = float(ws.cell(r, 9).value)
    top1_excel = float(ws.cell(r, 10).value)
    time_excel = float(ws.cell(r, 11).value)
    latency_excel = float(ws.cell(r, 12).value)
    
    run_dir = Path('runs') / exp_id / 'seed0'
    assert run_dir.exists(), f'Missing {run_dir}'
    
    with open(run_dir / 'summary.json', 'r', encoding='utf-8') as f:
        summary = json.load(f)
    with open(run_dir / 'config.json', 'r', encoding='utf-8') as f:
        cfg = json.load(f)
    hist = pd.read_csv(run_dir / 'history.csv')
    val_logits = np.load(run_dir / 'val_logits.npy')
    
    # Check shape
    assert val_logits.shape == (3501, 9), f'Bad shape {val_logits.shape}'
    probs = np.exp(val_logits - val_logits.max(axis=1, keepdims=True))
    probs /= probs.sum(axis=1, keepdims=True)
    preds = probs.argmax(axis=1)
    
    recalc_metrics = ev.compute_metrics(y_true_expected, preds, probs)
    recalc_f1 = round(float(recalc_metrics['macro_f1']), 4)
    recalc_top1 = round(float(recalc_metrics['top1']), 4)
    
    # Check prediction csv
    pred_csv = Path('predictions') / f'{exp_id}_seed0_val.csv'
    assert pred_csv.exists(), f'Missing {pred_csv}'
    pdf = pd.read_csv(pred_csv)
    assert len(pdf) == 3501, f'Bad pred csv len {len(pdf)}'
    
    # Check benchmark match
    b_item = next(item for item in bench if item['model'] == backbone)
    b_params = b_item['params_m']
    b_gmacs = b_item['gmacs']
    b_latency = b_item['p50']
    
    avg_train_time = round(hist['epoch_time_s'].mean(), 2)
    
    print(f'[{exp_id}: {backbone}]')
    print(f'  Params: Excel={params_excel:.3f}, Bench={b_params:.3f}')
    print(f'  GMACs:  Excel={gmac_excel:.3f}, Bench={b_gmacs:.3f}')
    print(f'  Latency: Excel={latency_excel:.2f}, Bench p50={b_latency:.2f}')
    print(f'  Time/ep: Excel={time_excel:.2f}, Hist mean={avg_train_time:.2f}')
    print(f'  Val F1: Excel={f1_excel:.4f}, Summary={summary["best_val_macro_f1"]:.4f}, Recalc={recalc_f1:.4f}')
    print(f'  Val Top1: Excel={top1_excel:.4f}, Summary={summary["best_val_top1"]:.4f}, Recalc={recalc_top1:.4f}')
    
    assert abs(params_excel - b_params) < 1e-3, 'Params mismatch'
    assert abs(gmac_excel - b_gmacs) < 1e-3, 'GMACs mismatch'
    assert abs(latency_excel - b_latency) < 1e-2, 'Latency mismatch'
    assert abs(time_excel - avg_train_time) < 0.1, 'Train time mismatch'
    assert abs(f1_excel - recalc_f1) < 1e-4, 'F1 mismatch'
    assert abs(top1_excel - recalc_top1) < 1e-4, 'Top1 mismatch'
    print(f'  ==> {exp_id} PASSED ALL INTEGRITY CHECKS!\n')

print('ALL 5 BACKBONES VERIFIED SUCCESSFULLY AGAINST RAW ARTIFACTS AND EVAL.PY!')
